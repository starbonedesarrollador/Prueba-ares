#!/usr/bin/env python3
"""
Generador de datos de demostración · Tornillería Oriente (DEMO — DATOS DE PRUEBA)

Crea, a partir de data/base/productos_base.csv (códigos, descripciones, categorías,
estantes y unidades del catálogo, más un costo de referencia FICTICIO):

  data/demo-db.json                 base de datos completa que carga la demo
  data/csv/*.csv                    los mismos datos en tablas planas (para Ares, Excel, etc.)
  data/listas-proveedores/*.xlsx    listas de precios de proveedores ficticios

Todo lo demás (existencias, costos, precios, clientes, proveedores, ventas, compras,
gastos, cobros, caja y tasas) es inventado. La semilla fija hace que el resultado sea
siempre igual: volver a correr el script restaura los datos originales de la demo.

Uso:
  python3 tools/generar_demo.py                 # fecha de referencia 2026-10-21
  python3 tools/generar_demo.py --fecha 2026-11-30 --semilla 7

Requiere Python 3.9+ y openpyxl (pip install openpyxl) para los .xlsx.
"""
import argparse, csv, json, math, os, random, datetime as dt
from collections import defaultdict

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(RAIZ, 'data')
TZ = dt.timedelta(hours=-4)  # Venezuela, UTC-4 todo el año

ap = argparse.ArgumentParser()
ap.add_argument('--fecha', default='2026-10-21', help='día "hoy" de la demo (AAAA-MM-DD)')
ap.add_argument('--semilla', type=int, default=20261021)
A = ap.parse_args()
R = random.Random(A.semilla)
HOY = dt.date.fromisoformat(A.fecha)

def ms(d, h=12, mi=0):
    x = dt.datetime(d.year, d.month, d.day, h, mi) - TZ
    return int((x - dt.datetime(1970, 1, 1)).total_seconds() * 1000)
def dk(d): return d.isoformat()
def r2(x): return round(x + 1e-9, 2)
def r4(x): return round(x + 1e-12, 4)

D1, D2, MARGEN, MINIMO = 0.53, 0.25, 3, 5   # reglas de precio del sistema (no son datos)

# ---------------------------------------------------------------- tasas BCV ficticias
dias = [HOY - dt.timedelta(days=i) for i in range(130, -1, -1)]
tasas, t = {}, 352.40
for d in dias:
    if d.weekday() < 5:
        t *= 1 + R.gauss(0.0011, 0.0009)
    tasas[d] = round(t, 2)
TASA_HOY = tasas[HOY]
def tasa(d): return tasas.get(d, TASA_HOY)

# ---------------------------------------------------------------- productos
base = list(csv.DictReader(open(os.path.join(DATA, 'base', 'productos_base.csv'), encoding='utf-8')))
prods = []
for b in base:
    c = float(b['costo_demo_base'])
    u = R.random()
    if u < 0.14:   exist = 0                                   # agotado
    elif u < 0.28: exist = R.randint(1, 4)                     # poca existencia (bajo el mínimo de 5)
    else:
        top = 2500 if c < 0.05 else 900 if c < 0.3 else 250 if c < 1.5 else 60
        exist = max(6, int(math.exp(R.uniform(math.log(6), math.log(top)))))
    p = {'pid': b['pid'], 'codigo': b['codigo'], 'desc': b['descripcion'], 'cat': b['categoria'], 'ubic': b['estante'],
         'unidad': b['unidad_venta'], 'costo': c, 'exist': exist, 'fijo': None, 'minimo': None}
    if b['usa_precio_fijo'] == '1':
        p['fijo'] = r2(max(0.25, c * (1 - D1) * (1 + MARGEN) * R.uniform(0.95, 1.25)))
    if R.random() < 0.02:
        p['costo'] = None                                      # algunos sin costo: se venden con precio a mano
    if exist > 0 and R.random() < 0.06:
        p['minimo'] = R.choice([10, 20, 50])
    prods.append(p)
byCode = {p['codigo']: p for p in prods}
def H(p): return None if p['costo'] is None else p['costo'] * (1 - D1)
def J(p): return p['fijo'] if p['fijo'] is not None else (None if H(p) is None else H(p) * (1 + MARGEN))
vendibles = [p for p in prods if p['exist'] > 0 and J(p)]
R.shuffle(vendibles)
populares = vendibles[:1500]
pesos = [1 / (i + 40) for i in range(len(populares))]

def cant_para(p, mayor=False):
    j = J(p)
    if j < 0.08: q = R.choice([10, 12, 20, 25, 50, 100])
    elif j < 0.5: q = R.choice([2, 4, 6, 10, 12, 20])
    elif j < 3: q = R.choice([1, 2, 2, 4, 6])
    else: q = R.choice([1, 1, 2])
    return q * (R.choice([3, 5, 10]) if mayor else 1)

# ---------------------------------------------------------------- clientes ficticios
NOMBRES = ['Taller Mecánico El Pistón, C.A.', 'Repuestos La Bujía, C.A.', 'Ferretería Punta Brava, C.A.', 'Inversiones Tuerca Fina, C.A.',
  'Servicios Hidráulicos Morichal, C.A.', 'Construcciones Cerro Alto, C.A.', 'Autopartes Bahía Azul, C.A.', 'Metalúrgica Los Samanes, C.A.',
  'Taller Diésel Puerto Viejo', 'Carpintería Metálica El Yunque', 'Agropecuaria La Ceiba, C.A.', 'Mantenimiento Industrial Delta, C.A.',
  'Ana Rivas', 'Carlos Mendoza', 'Luisa Fermín', 'José Salazar', 'María Carvajal', 'Pedro Lugo', 'Rosa Marcano', 'Juan Bastidas',
  'Andrés Velásquez', 'Gabriela Rondón', 'Miguel Guzmán', 'Daniela Rojas', 'Rafael Patiño', 'Carmen Silva', 'Luis Acosta', 'Elena Brito',
  'Héctor Figuera', 'Paola Millán', 'Taller Las Gaviotas', 'Moto Repuestos El Rayo', 'Herrería San Rafael', 'Electromecánica Oriente Sur, C.A.',
  'Refrigeración Polar Demo, C.A.', 'Transporte Costa Firme, C.A.', 'Yolanda Campos', 'Oscar Natera', 'Ramón Hernández', 'Beatriz Aular']
clientes = {}
for i, n in enumerate(NOMBRES):
    emp = ',' in n or n.split()[0] in ('Taller', 'Moto', 'Herrería', 'Carpintería')
    key = ('J99%07d' % (i + 101)) if emp else ('V99%06d' % (i + 101))      # documentos imposibles: rango 99… no asignado
    doc = ('J-99%07d-0' % (i + 101)) if emp else ('V-99.%03d.%03d' % ((i + 101) // 1000, (i + 101) % 1000))
    clientes[key] = {'nombre': n.upper(), 'doc': doc, 'tel': '', 'dir': '', 'notas': 'Cliente ficticio de la demo', 'creado': ms(HOY - dt.timedelta(days=200)), 'origen': 'demo'}
    if emp and R.random() < 0.25:
        clientes[key]['especial'] = True; clientes[key]['retPct'] = 0.75
CK = list(clientes)

# ---------------------------------------------------------------- historial (formato de la hoja: ventas de 90 días)
hist = defaultdict(lambda: {'lineas': [], 'nombres': [], 'descs': {}})
for i in range(92, -1, -1):
    d = HOY - dt.timedelta(days=i)
    if d.weekday() == 6: continue
    n = R.randint(30, 60) if d.weekday() < 5 else R.randint(18, 34)
    if d == HOY: n = n // 2                                        # el día de hoy va por la mitad
    mk = d.strftime('%Y-%m'); doc = hist[mk]
    for _ in range(n):
        p = R.choices(populares, pesos)[0]
        nombrado = R.random() < 0.3
        q = cant_para(p, mayor=nombrado and R.random() < 0.3)
        usd = r2(q * J(p)); ts = tasa(d)
        pm = efbs = efusd = punto = 0.0
        f = R.random()
        if f < 0.38: pm = r2(usd * ts)
        elif f < 0.62: punto = r2(usd * ts)
        elif f < 0.72: efbs = r2(usd * ts)
        else: efusd = usd
        ni = -1
        if nombrado:
            nom = clientes[R.choice(CK)]['nombre']
            if nom not in doc['nombres']: doc['nombres'].append(nom)
            ni = doc['nombres'].index(nom)
        doc['lineas'].append([d.day, p['codigo'], q, usd, r2(usd * ts), ts, pm, efbs, efusd, punto, ni])
historial = {}
for mk, doc in hist.items():
    doc['mes'] = mk; historial[mk] = doc

# ---------------------------------------------------------------- proveedores y listas (unidad y caja)
proveedores = {
  'tornorte':   {'nombre': 'DISTRIBUIDORA TORNILLOS DEL NORTE, C.A. (DEMO)', 'corto': 'Tornorte', 'descBs': 0.40, 'descExtra': 0.20, 'pref': 'TN'},
  'pernosandes':{'nombre': 'IMPORTADORA PERNOS LOS ANDES, C.A. (DEMO)', 'corto': 'Pernos Andes', 'descBs': 0.0, 'descExtra': 0.25, 'pref': 'PA'},
  'ferrecaribe':{'nombre': 'FERRETERA INDUSTRIAL CARIBE, C.A. (DEMO)', 'corto': 'FerreCaribe', 'descBs': 0.30, 'descExtra': 0.15, 'pref': 'FC'},
  'autotuercas':{'nombre': 'AUTOTUERCAS DEL ORIENTE, C.A. (DEMO)', 'corto': 'AutoTuercas', 'descBs': 0.50, 'descExtra': 0.25, 'pref': 'AT'},
}
con_costo = [p for p in prods if p['costo']]
R.shuffle(con_costo)
comunes = con_costo[:160]                          # están en las tres listas
listas, lineas_xlsx = {}, {}
for li, pid in enumerate(['tornorte', 'pernosandes', 'ferrecaribe']):
    pv = proveedores[pid]
    propios = con_costo[160 + li * 140: 160 + (li + 1) * 140]
    items, filas = [], []
    for k, p in enumerate(comunes + propios):
        unit = r4(p['costo'] * R.uniform(0.88, 1.12))     # precio de lista cerca del costo de lista que tiene el producto
        cod = '%s-%05d' % (pv['pref'], 1000 + k * 7 + li)
        chica = unit < 0.6
        caja_n = R.choice([25, 50, 100]) if unit < 0.15 else R.choice([10, 20, 25, 50])
        solo_caja = chica and R.random() < 0.25
        fecha = ms(HOY - dt.timedelta(days=R.randint(5, 25)))
        if not solo_caja:
            items.append([cod, p['desc'], unit, p['codigo'], 0, '', 1, 'UNIDAD', 1, unit])
            filas.append([cod, p['desc'], 'UNIDAD', 1, unit, unit, p['codigo']])
        if chica:
            pu = r4(unit * R.uniform(0.85, 0.94))     # en caja sale más barato por unidad
            pc = r2(pu * caja_n)
            items.append([cod + '-C' + str(caja_n), p['desc'] + ' · CAJA x' + str(caja_n), pu, p['codigo'], 0, '', 1, 'CAJA', caja_n, pc])
            filas.append([cod + '-C' + str(caja_n), p['desc'], 'CAJA', caja_n, pc, pu, p['codigo']])
    lid = pid + '-general'
    listas[lid] = {'prov': pid, 'nombre': pv['corto'] + ' · Lista general (demo)', 'fecha': ms(HOY - dt.timedelta(days=12)),
                   'cols': ['codProv', 'desc', 'precioUnidad', 'nuestroCodigo', 'oferta', 'sugerido', 'confianza', 'presentacion', 'unidadesPorEmpaque', 'precioEmpaque'],
                   'items': items}
    lineas_xlsx[pid] = filas

# ---------------------------------------------------------------- cuentas
cuentas = {
  'banco_pm':    {'nombre': 'Banco Demo · Pago móvil', 'banco': 'Banco Demo', 'tipo': 'pago_movil', 'mon': 'BS', 'datos': '', 'titular': 'Tornillería Oriente (demo)', 'activa': True},
  'banco_emp':   {'nombre': 'Banco Demo · Punto y transferencias', 'banco': 'Banco Demo', 'tipo': 'punto', 'mon': 'BS', 'datos': '', 'titular': 'Tornillería Oriente (demo)', 'activa': True},
  'zelle_demo':  {'nombre': 'Zelle (demo)', 'banco': 'Zelle', 'tipo': 'zelle', 'mon': 'USD', 'datos': '', 'titular': '', 'activa': True},
  'binance_demo':{'nombre': 'Binance (demo)', 'banco': 'Binance', 'tipo': 'otro_usd', 'mon': 'USD', 'datos': '', 'titular': '', 'activa': True},
  'efectivo_usd':{'nombre': 'Efectivo $', 'banco': 'Efectivo', 'tipo': 'efectivo', 'mon': 'USD', 'datos': '', 'titular': '', 'activa': True, 'fijo': True},
  'efectivo_bs': {'nombre': 'Efectivo Bs', 'banco': 'Efectivo', 'tipo': 'efectivo', 'mon': 'BS', 'datos': '', 'titular': '', 'activa': True, 'fijo': True},
}
t_saldo = ms(HOY - dt.timedelta(days=30), 7, 30)
for k, s in [('banco_pm', 18500.0), ('banco_emp', 42300.0), ('zelle_demo', 380.0), ('binance_demo', 210.0), ('efectivo_usd', 150.0), ('efectivo_bs', 6200.0)]:
    cuentas[k].update({'creado': t_saldo, 'por': 'demo_admin', 'saldoIni': s, 'saldoAt': t_saldo})
CTA = {k: v['nombre'] for k, v in cuentas.items()}
def pago(key, usd_, d):
    c = cuentas[key]; ts = tasa(d)
    monto = r2(usd_) if c['mon'] == 'USD' else r2(usd_ * ts)
    return {'key': 'c_' + key if not c.get('fijo') else key, 'label': c['nombre'], 'mon': c['mon'], 'monto': monto, 'usd': r2(usd_), 'tasa': ts}

# ---------------------------------------------------------------- ventas del sistema (últimos 9 días)
ventas, vendidos = {}, defaultdict(float)
nro = 0
candidatos = [p for p in populares if p['exist'] >= 30 and p['costo']]
for i in range(9, -1, -1):
    d = HOY - dt.timedelta(days=i)
    if d.weekday() == 6: continue
    for _ in range(R.randint(1, 4) if d != HOY else 2):
        nro += 1
        hh = R.randint(8, 16) if d != HOY else R.choice([8, 9])
        t = ms(d, hh, R.randint(0, 59)); ts = tasa(d)
        modo = 'usd' if R.random() < 0.3 else 'bs'
        credito = modo == 'bs' and R.random() < 0.22
        con_iva = not credito and R.random() < 0.15
        ck = R.choice(CK) if (credito or con_iva or R.random() < 0.35) else None
        items, sub = [], 0.0
        for p in R.sample(candidatos, R.randint(1, 4)):
            q = cant_para(p)
            precio = r4(J(p) * (1 - D2)) if modo == 'usd' else r4(J(p))
            tot = r2(q * precio); sub += tot; vendidos[p['codigo']] += q
            items.append({'id': p['pid'], 'codigo': p['codigo'], 'desc': p['desc'], 'ubic': p['ubic'], 'cant': q, 'precio': precio, 'costo': r4(H(p)), 'total': tot})
        sub = r2(sub); iva = r2(sub * 0.16) if con_iva else 0.0; total = r2(sub + iva)
        if credito:
            pg = [{'metodo': 'credito', 'cuenta': None, 'cuentaNombre': '', 'mon': 'USD', 'monto': total, 'ref': ''}]
        elif modo == 'usd':
            k = R.choice(['efectivo_usd', 'zelle_demo'])
            pg = [{'metodo': 'efectivo_usd' if k == 'efectivo_usd' else 'zelle', 'cuenta': None if k == 'efectivo_usd' else k, 'cuentaNombre': '' if k == 'efectivo_usd' else CTA[k], 'mon': 'USD', 'monto': total, 'ref': ''}]
        else:
            k = R.choice(['banco_pm', 'banco_pm', 'banco_emp', 'efectivo_bs'])
            met = {'banco_pm': 'pago_movil', 'banco_emp': 'punto', 'efectivo_bs': 'efectivo_bs'}[k]
            pg = [{'metodo': met, 'cuenta': None if k == 'efectivo_bs' else k, 'cuentaNombre': '' if k == 'efectivo_bs' else CTA[k], 'mon': 'BS', 'monto': r2(total * ts), 'ref': '' if k == 'efectivo_bs' else str(R.randint(100000, 999999))}]
        c = clientes.get(ck) if ck else None
        v = {'n': 'V-%06d' % nro, 'nro': nro, 'cod': 'V-%06d' % nro, 't': t, 'dia': dk(d), 'registrada': t, 'baseImponible': sub, 'tasa': ts, 'subtotal': sub,
             'iva': iva, 'ivaTasa': 0.16 if iva else 0, 'igtf': 0, 'igtfTasa': 0, 'total': total, 'totalBs': r2(total * ts),
             'cliente': c['nombre'] if c else '', 'clienteId': ck, 'clienteDoc': c['doc'] if c else '', 'clienteTel': '', 'por': 'demo_admin',
             'anulada': False, 'items': items, 'pagos': pg, 'modo': modo, 'credito': total if credito else 0, 'vuelto': 0, 'vueltoMon': None}
        if c and c.get('especial') and iva:
            ret = r2(iva * c['retPct']); v['retencion'] = {'pct': c['retPct'], 'monto': ret, 'montoBs': r2(ret * ts), 'comprobante': '', 'estado': 'pendiente'}
            v['neto'] = r2(total - ret); v['netoBs'] = r2(v['neto'] * ts)
        ventas['v%04d' % nro] = v
# una venta anulada (la mercancía vuelve)
anul = [k for k, v in ventas.items() if not v['credito']][2]
ventas[anul].update({'anulada': True, 'anuladaAt': ventas[anul]['t'] + 600000, 'anuladaPor': 'demo_admin'})
for it in ventas[anul]['items']: vendidos[it['codigo']] -= it['cant']
# la existencia guardada ya incluye lo vendido en el sistema (el sistema la descuenta al mostrar)
for code, q in vendidos.items(): byCode[code]['exist'] += int(q)

# ---------------------------------------------------------------- abonos a ventas a crédito
abonos, na = {}, 0
cred = [(k, v) for k, v in ventas.items() if v['credito'] and not v['anulada']]
for k, v in cred[:3]:
    na += 1; d = dt.date.fromisoformat(v['dia']) + dt.timedelta(days=1)
    if d > HOY: d = HOY
    monto = r2(v['credito'] * R.choice([0.3, 0.5, 1.0]))
    p = pago('banco_pm', monto, d); p['ref'] = str(R.randint(100000, 999999))
    abonos['ab%03d' % na] = {'tipo': 'cxc', 'ref': k, 't': ms(d, 15), 'dia': dk(d), 'pago': p, 'quien': v['cliente'], 'por': 'demo_admin', 'nro': na, 'n': 'R-%06d' % na}

# ---------------------------------------------------------------- compras a proveedores
compras, nc = {}, 0
for i, (pid, dias_atras, cond, en) in enumerate([('tornorte', 27, 'contado', 'bs'), ('pernosandes', 21, 'credito', 'usd'), ('ferrecaribe', 16, 'contado', 'usd'),
                                                  ('tornorte', 11, 'credito', 'bs'), ('autotuercas', 6, 'contado', 'bs'), ('ferrecaribe', 2, 'credito', 'bs')]):
    nc += 1; pv = proveedores[pid]; d = HOY - dt.timedelta(days=dias_atras)
    lista = [x for x in listas.get(pid + '-general', {'items': []})['items'] if x[7] == 'UNIDAD'] or [[None, p['desc'], r4(p['costo'] * R.uniform(0.9, 1.1)), p['codigo'], 0, '', 1, 'UNIDAD', 1, 0] for p in R.sample(con_costo, 12)]
    sel = R.sample(lista, min(len(lista), R.randint(4, 9)))
    its, total = [], 0.0
    for x in sel:
        p = byCode[x[3]]; q = R.choice([50, 100, 200, 500]) if x[2] < 0.3 else R.choice([10, 20, 30, 50])
        real = x[2] * (1 - pv['descBs']) * ((1 - pv['descExtra']) if en == 'usd' else 1)
        its.append({'id': p['pid'], 'codigo': p['codigo'], 'desc': p['desc'], 'cant': q, 'G': x[2], 'costoReal': r4(real)}); total += real * q
        p['ultCompra'] = {'prov': pid, 'G': x[2], 'descBs': pv['descBs'], 'descExtra': pv['descExtra'], 'dia': dk(d)}
    total = r2(total)
    compras['cp%03d' % nc] = {'t': ms(d, 10), 'dia': dk(d), 'proveedorId': pid, 'proveedorNombre': pv['nombre'], 'factura': 'DEMO-%04d' % (2100 + nc * 13),
        'pagoEn': en, 'condicion': cond, 'descBs': pv['descBs'], 'descExtra': pv['descExtra'], 'total': total,
        'pago': (pago('banco_emp' if en == 'bs' else 'efectivo_usd', total, d) if cond == 'contado' else None), 'por': 'demo_admin', 'anulada': False,
        'items': its, 'nro': nc, 'n': 'C-%06d' % nc}
# un pago a proveedor
c2 = compras['cp002']; na += 1
abonos['ab%03d' % na] = {'tipo': 'cxp', 'ref': 'cp002', 't': ms(HOY - dt.timedelta(days=9), 11), 'dia': dk(HOY - dt.timedelta(days=9)),
    'pago': dict(pago('efectivo_usd', r2(c2['total'] * 0.4), HOY - dt.timedelta(days=9)), ref=''), 'quien': c2['proveedorNombre'], 'por': 'demo_admin', 'nro': na, 'n': 'R-%06d' % na}

# ---------------------------------------------------------------- gastos
gastos, ng = {}, 0
for dias_atras, cat, desc, usd_, key in [(29, 'ALQUILER', 'Alquiler del local (demo)', 450, 'banco_emp'), (26, 'SERVICIOS', 'Electricidad (demo)', 38.5, 'banco_pm'),
        (24, 'TRANSPORTE', 'Flete de mercancía (demo)', 25, 'efectivo_usd'), (20, 'NÓMINA', 'Pago quincenal vendedor (demo)', 160, 'efectivo_usd'),
        (17, 'MANTENIMIENTO', 'Reparación de estantería (demo)', 42, 'efectivo_bs'), (14, 'SERVICIOS', 'Internet (demo)', 30, 'banco_pm'),
        (12, 'PAPELERÍA', 'Talonarios y bolsas (demo)', 12.5, 'efectivo_bs'), (8, 'TRANSPORTE', 'Gasolina (demo)', 18, 'efectivo_usd'),
        (5, 'NÓMINA', 'Pago quincenal vendedor (demo)', 160, 'efectivo_usd'), (3, 'SERVICIOS', 'Agua (demo)', 9.5, 'banco_pm'),
        (1, 'LIMPIEZA', 'Artículos de limpieza (demo)', 7.8, 'efectivo_bs')]:
    ng += 1; d = HOY - dt.timedelta(days=dias_atras)
    gastos['g%03d' % ng] = {'t': ms(d, 16), 'dia': dk(d), 'categoria': cat, 'desc': desc, 'pago': pago(key, usd_, d), 'por': 'demo_admin', 'nro': ng, 'n': 'G-%06d' % ng}

# ---------------------------------------------------------------- compra de divisas
movcta, nm = {}, 0
for dias_atras, usd_, sobre in [(25, 200, 1.04), (13, 150, 1.05), (4, 120, 1.035)]:
    nm += 1; d = HOY - dt.timedelta(days=dias_atras); ts = tasa(d); pag = round(ts * sobre, 2)
    movcta['m%03d' % nm] = {'t': ms(d, 13), 'dia': dk(d), 'tipo': 'compra', 'a': 'banco_emp', 'aNombre': CTA['banco_emp'], 'aMon': 'BS', 'montoA': r2(usd_ * pag),
        'key': 'efectivo_usd', 'nombre': 'Efectivo $', 'mon': 'USD', 'monto': usd_, 'tasa': pag, 'tasaBcv': ts, 'nota': 'Compra de divisas (demo)', 'txt': '', 'anulado': False, 'nro': nm, 'n': 'M-%06d' % nm}

# ---------------------------------------------------------------- saldos iniciales: que cada cuenta termine en un monto realista
meta_saldo = {'banco_pm': 96000.0, 'banco_emp': 134000.0, 'zelle_demo': 420.0, 'binance_demo': 210.0, 'efectivo_usd': 385.0, 'efectivo_bs': 9200.0}
neto = defaultdict(float)
def cta_de(pg):
    if pg.get('cuenta'): return pg['cuenta']
    return {'efectivo_usd': 'efectivo_usd', 'efectivo_bs': 'efectivo_bs'}.get(pg.get('metodo') or pg.get('key'))
for v in ventas.values():
    if v['anulada']: continue
    for pg in v['pagos']:
        k = cta_de(pg)
        if k: neto[k] += pg['monto']
for a_ in abonos.values():
    k = a_['pago']['key'].replace('c_', ''); neto[k] += a_['pago']['monto'] * (1 if a_['tipo'] == 'cxc' else -1)
for g in gastos.values():
    k = g['pago']['key'].replace('c_', ''); neto[k] -= g['pago']['monto']
for c_ in compras.values():
    if c_['pago']: k = c_['pago']['key'].replace('c_', ''); neto[k] -= c_['pago']['monto']
for m_ in movcta.values():
    neto[m_['a']] -= m_['montoA']; neto[m_['key']] += m_['monto']
for k, tgt in meta_saldo.items():
    cuentas[k]['saldoIni'] = r2(tgt - neto[k])

# ---------------------------------------------------------------- caja: días anteriores cerrados, hoy abierta
cajas = {}
for i in range(5, -1, -1):
    d = HOY - dt.timedelta(days=i)
    if d.weekday() == 6: continue
    vd = [v for v in ventas.values() if v['dia'] == dk(d) and not v['anulada']]
    c = {'dia': dk(d), 'apertura': {'usd': 40, 'bs': 2000, 'at': ms(d, 8, 5), 'por': 'demo_admin'}}
    if d != HOY:
        tot = r2(sum(v['total'] for v in vd))
        c['cierre'] = {'at': ms(d, 17, 40), 'por': 'demo_admin', 'esperado': {}, 'labels': {}, 'contado': {}, 'diferencias': {}, 'notas': 'Cierre de prueba (demo)',
                       'ventas': len(vd), 'total': tot, 'totalBs': r2(tot * tasa(d))}
    cajas[dk(d)] = c

# ---------------------------------------------------------------- deudas iniciales (traídas "de la hoja")
items_d = []
for i, ck in enumerate(R.sample(CK, 10)):
    d = HOY - dt.timedelta(days=R.randint(8, 140))
    det, tot = [], 0.0
    for p in R.sample(populares, R.randint(1, 4)):
        q = cant_para(p, True); tt = r2(q * J(p)); tot += tt
        det.append({'c': p['codigo'], 'd': p['desc'], 'q': q, 'u': r4(J(p)), 't': tt})
    items_d.append({'id': 'c%d' % (i + 1), 'tipo': 'cxc', 't': ms(d), 'quien': clientes[ck]['nombre'], 'ref': 'Hoja de cobros (demo) fila %d' % (40 + i * 7), 'total': r2(tot), 'detalle': det})
for i, (pid, mto, dd) in enumerate([('autotuercas', 612.40, 33), ('pernosandes', 288.15, 18), ('tornorte', 940.00, 9)]):
    items_d.append({'id': 'p%d' % (i + 1), 'tipo': 'cxp', 't': ms(HOY - dt.timedelta(days=dd)), 'quien': proveedores[pid]['nombre'], 'ref': 'Factura DEMO-%04d' % (900 + i * 37), 'total': mto,
                    'detalle': [{'c': '', 'd': 'Factura de prueba', 'q': 1, 'u': mto, 't': mto}]})

# ---------------------------------------------------------------- armar la base completa
catalogo = defaultdict(dict)
for idx, p in enumerate(prods):
    it = {'codigo': p['codigo'], 'desc': p['desc'], 'cat': p['cat'], 'ubic': p['ubic'], 'unidad': p['unidad'], 'exist': p['exist'],
          'costo': p['costo'], 'd1': None, 'd2': None, 'base': 'H', 'fijo': p['fijo'], 'minimo': p['minimo']}
    if p.get('ultCompra'): it['ultCompra'] = p['ultCompra']
    catalogo['c%02d' % (idx // 100)][p['pid']] = it
t_tasa = ms(HOY, 8, 5)
tasas_docs = {'t%s' % d.strftime('%Y%m%d'): {'tasa': tasas[d], 'anterior': tasas[d - dt.timedelta(days=1)] if (d - dt.timedelta(days=1)) in tasas else None,
              'at': ms(d, 8, 5), 'por': 'auto', 'nota': 'Tasa ficticia de la demo'} for d in dias[-45:] if d.weekday() < 5}
col = {
  'config': {'general': {'d1': D1, 'd2': D2, 'margen': MARGEN, 'minimo': MINIMO, 'tasa': TASA_HOY, 'tasaAt': t_tasa, 'tasaPor': 'auto',
             'empresa': {'razon': 'TORNILLERÍA ORIENTE · DEMO', 'rif': 'J-00000000-0', 'tel': '', 'dir': 'Dirección de prueba (demo)', 'correo': '', 'validez': 3,
                         'pie': 'DEMO — DATOS DE PRUEBA. Documento sin validez.', 'retPct': 0.75}}},
  'catalogo': {k: {'items': v} for k, v in catalogo.items()},
  'tasas': tasas_docs, 'cambios': {}, 'avisos': {'importacion': {'lista': [], 'revisados': {}}},
  'clientes': clientes, 'cuentas': cuentas, 'proveedores': {k: {kk: vv for kk, vv in v.items() if kk != 'pref'} | {'creado': ms(HOY - dt.timedelta(days=300))} for k, v in proveedores.items()},
  'listas': listas, 'cruces': {}, 'ventas': ventas, 'compras': compras, 'gastos': gastos, 'abonos': abonos, 'movcta': movcta,
  'presupuestos': {}, 'traspasos': {}, 'cajas': cajas, 'deudas': {'hoja': {'items': items_d}}, 'historial': historial,
  'correlativos': {'venta': {'ultimo': nro}, 'compra': {'ultimo': nc}, 'gasto': {'ultimo': ng}, 'abono': {'ultimo': na}, 'movcta': {'ultimo': nm}, 'presupuesto': {'ultimo': 0}, 'traspaso': {'ultimo': 0}},
}
meta = {'demo': True, 'aviso': 'DEMO — DATOS DE PRUEBA', 'fechaReferencia': dk(HOY), 'semilla': A.semilla, 'version': 1,
        'nota': 'Al cargar, la demo corre todas las fechas para que la fecha de referencia sea el día de hoy.'}
os.makedirs(DATA, exist_ok=True)
json.dump({'meta': meta, 'colecciones': col}, open(os.path.join(DATA, 'demo-db.json'), 'w', encoding='utf-8'), ensure_ascii=False, separators=(',', ':'))

# ---------------------------------------------------------------- una carpeta JSON por colección (misma información, más fácil de leer)
JD = os.path.join(DATA, 'json'); os.makedirs(JD, exist_ok=True)
for k, v in col.items():
    json.dump(v, open(os.path.join(JD, k + '.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
json.dump(meta, open(os.path.join(JD, '_meta.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)

# ---------------------------------------------------------------- CSV planos
C = os.path.join(DATA, 'csv'); os.makedirs(C, exist_ok=True)
def wcsv(n, head, rows):
    with open(os.path.join(C, n), 'w', newline='', encoding='utf-8-sig') as fh:
        w = csv.writer(fh); w.writerow(head); w.writerows(rows)
def est(p):
    m = p['minimo'] or MINIMO
    return 'agotado' if p['exist'] - vendidos.get(p['codigo'], 0) <= 0 else 'poca existencia' if p['exist'] - vendidos.get(p['codigo'], 0) <= m else 'disponible'
wcsv('productos.csv', ['codigo', 'descripcion', 'categoria', 'estante', 'unidad_venta', 'existencia', 'estado', 'costo_lista_usd', 'costo_pagando_bs_usd', 'costo_pagando_divisas_usd',
     'precio_venta_bcv_usd', 'precio_venta_divisas_usd', 'precio_fijo', 'existencia_minima'],
     [[p['codigo'], p['desc'], p['cat'], p['ubic'], p['unidad'] or 'UNIDAD', int(p['exist'] - vendidos.get(p['codigo'], 0)), est(p), p['costo'] if p['costo'] else '',
       r4(H(p)) if H(p) else '', r4(H(p) * (1 - D2)) if H(p) else '', r4(J(p)) if J(p) else '', r4(J(p) * (1 - D2)) if J(p) else '',
       'sí' if p['fijo'] is not None else 'no', p['minimo'] or MINIMO] for p in prods])
wcsv('clientes.csv', ['clave', 'documento_ficticio', 'nombre', 'contribuyente_especial', 'telefono', 'direccion'],
     [[k, c['doc'], c['nombre'], 'sí' if c.get('especial') else 'no', '', ''] for k, c in clientes.items()])
wcsv('proveedores.csv', ['id', 'nombre', 'nombre_corto', 'descuento_pagando_bs', 'descuento_extra_pagando_divisas'],
     [[k, v['nombre'], v['corto'], v['descBs'], v['descExtra']] for k, v in proveedores.items()])
wcsv('listas_precios.csv', ['proveedor', 'codigo_proveedor', 'descripcion', 'presentacion', 'unidades_por_empaque', 'precio_empaque_usd', 'precio_por_unidad_usd', 'nuestro_codigo'],
     [[pid] + f[:3] + [f[3], f[4], f[5], f[6]] for pid, F in lineas_xlsx.items() for f in F])
mult = defaultdict(set)
for pid, F in lineas_xlsx.items():
    for f in F: mult[f[6]].add(pid)
wcsv('productos_en_varios_proveedores.csv', ['nuestro_codigo', 'descripcion', 'proveedores'],
     [[c, byCode[c]['desc'], ' / '.join(sorted(proveedores[x]['corto'] for x in s))] for c, s in sorted(mult.items()) if len(s) > 1])
filas = []
for mk, doc in sorted(historial.items()):
    y, m = map(int, mk.split('-'))
    for l in doc['lineas']:
        filas.append([dk(dt.date(y, m, l[0])), l[1], byCode[l[1]]['desc'], l[2], l[3], l[4], l[5], doc['nombres'][l[10]] if l[10] >= 0 else '',
                      'Pago móvil' if l[6] else 'Punto' if l[9] else 'Efectivo Bs' if l[7] else 'Efectivo $'])
wcsv('historial_ventas_hoja.csv', ['fecha', 'codigo', 'descripcion', 'cantidad', 'total_usd', 'total_bs', 'tasa', 'cliente', 'forma_de_pago'], filas)
wcsv('ventas_sistema.csv', ['numero', 'fecha', 'cliente', 'modo', 'subtotal_usd', 'iva_usd', 'total_usd', 'total_bs', 'tasa', 'credito_usd', 'anulada', 'forma_de_pago'],
     [[v['n'], v['dia'], v['cliente'], v['modo'], v['subtotal'], v['iva'], v['total'], v['totalBs'], v['tasa'], v['credito'], 'sí' if v['anulada'] else 'no',
       (v['pagos'][0]['cuentaNombre'] or v['pagos'][0]['metodo'])] for v in ventas.values()])
wcsv('ventas_sistema_lineas.csv', ['numero', 'codigo', 'descripcion', 'cantidad', 'precio_usd', 'total_usd', 'costo_unitario_usd'],
     [[v['n'], it['codigo'], it['desc'], it['cant'], it['precio'], it['total'], it['costo']] for v in ventas.values() for it in v['items']])
wcsv('compras.csv', ['numero', 'fecha', 'proveedor', 'factura', 'condicion', 'pagado_en', 'total_usd', 'codigo', 'descripcion', 'cantidad', 'precio_lista_usd', 'costo_real_usd'],
     [[c['n'], c['dia'], c['proveedorNombre'], c['factura'], c['condicion'], c['pagoEn'], c['total'], it['codigo'], it['desc'], it['cant'], it['G'], it['costoReal']] for c in compras.values() for it in c['items']])
wcsv('gastos.csv', ['numero', 'fecha', 'categoria', 'descripcion', 'cuenta', 'moneda', 'monto', 'equivalente_usd'],
     [[g['n'], g['dia'], g['categoria'], g['desc'], g['pago']['label'], g['pago']['mon'], g['pago']['monto'], g['pago']['usd']] for g in gastos.values()])
wcsv('cobros_y_pagos.csv', ['recibo', 'fecha', 'tipo', 'quien', 'referencia', 'cuenta', 'moneda', 'monto', 'equivalente_usd'],
     [[a['n'], a['dia'], 'cobro a cliente' if a['tipo'] == 'cxc' else 'pago a proveedor', a['quien'], (ventas.get(a['ref']) or compras.get(a['ref']) or {}).get('n', a['ref']),
       a['pago']['label'], a['pago']['mon'], a['pago']['monto'], a['pago']['usd']] for a in abonos.values()])
wcsv('deudas_iniciales.csv', ['tipo', 'quien', 'fecha', 'referencia', 'total_usd'],
     [['por cobrar' if x['tipo'] == 'cxc' else 'por pagar', x['quien'], dk((dt.datetime(1970, 1, 1) + dt.timedelta(milliseconds=x['t']) + TZ).date()), x['ref'], x['total']] for x in items_d])
wcsv('movimientos_divisas.csv', ['numero', 'fecha', 'sale_de', 'monto_bs', 'entra_a', 'monto_usd', 'tasa_pagada', 'tasa_bcv'],
     [[m['n'], m['dia'], m['aNombre'], m['montoA'], m['nombre'], m['monto'], m['tasa'], m['tasaBcv']] for m in movcta.values()])
wcsv('tasas_bcv.csv', ['fecha', 'tasa'], [[dk(d), tasas[d]] for d in dias if d.weekday() < 5])
wcsv('cuentas.csv', ['id', 'nombre', 'moneda', 'saldo_inicial', 'fecha_saldo_inicial'], [[k, c['nombre'], c['mon'], c['saldoIni'], dk(HOY - dt.timedelta(days=30))] for k, c in cuentas.items()])

# ---------------------------------------------------------------- listas en Excel
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    X = os.path.join(DATA, 'listas-proveedores'); os.makedirs(X, exist_ok=True)
    for pid, F in lineas_xlsx.items():
        pv = proveedores[pid]; wb = Workbook(); ws = wb.active; ws.title = 'Lista de precios'
        ws.append(['DEMO — DATOS DE PRUEBA · ' + pv['nombre']]); ws['A1'].font = Font(bold=True, size=13)
        ws.append(['Precios ficticios en dólares. Descuento pagando en Bs: %d %% · extra pagando en divisas: %d %%' % (pv['descBs'] * 100, pv['descExtra'] * 100)])
        ws.append([])
        ws.append(['Código proveedor', 'Descripción', 'Presentación', 'Unidades por empaque', 'Precio por empaque $', 'Precio por unidad $', 'Nuestro código'])
        for c in ws[4]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='353535'); c.alignment = Alignment(vertical='center')
        for f in F: ws.append(f)
        for col_, w in zip('ABCDEFG', [18, 58, 14, 12, 16, 16, 18]): ws.column_dimensions[col_].width = w
        for row in ws.iter_rows(min_row=5, min_col=5, max_col=6):
            for c in row: c.number_format = '0.0000'
        ws.freeze_panes = 'A5'
        wb.save(os.path.join(X, 'lista_%s_demo.xlsx' % pid))
except ImportError:
    print('Aviso: falta openpyxl; no se generaron los .xlsx (pip install openpyxl).')

print('Listo:', len(prods), 'productos ·', len(clientes), 'clientes ·', sum(len(d['lineas']) for d in historial.values()), 'líneas de historial ·',
      len(ventas), 'ventas del sistema ·', len(compras), 'compras ·', sum(len(l['items']) for l in listas.values()), 'líneas en listas')
